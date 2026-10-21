import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DATA_PATH = r"C:\Users\User\AppData\Local\Temp\claude\C--Users-User-Desktop-tw-data-updater\1cb3a135-d2a6-4f7e-b9d2-604385fefaf2\scratchpad\currency_export_data.json"
OUT_PATH = r"C:\Users\User\Desktop\幣別比對清單_20261021.xlsx"

with open(DATA_PATH, encoding="utf-8") as f:
    data = json.load(f)

FONT = "Arial"
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
BODY_FONT = Font(name=FONT, size=10)
CONFIRM_FONT = Font(name=FONT, size=10, color="9C5700")
CONFIRM_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

HEADERS = [
    "基金代碼", "基金完整名稱（資料庫）", "基金完整名稱（官方資料）",
    "資料庫目前的幣別", "官方資料的幣別（原文）", "準備改成的ISO代碼",
    "官方資料來源與資料日期", "是否為避險級別", "名稱中是否出現多個幣別", "信心程度",
]
COL_WIDTHS = [16, 42, 42, 12, 14, 12, 26, 12, 14, 14]

def write_sheet(ws, rows):
    ws.sheet_view.showGridLines = False
    for col_idx, (header, width) in enumerate(zip(HEADERS, COL_WIDTHS), start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 28

    keys = ["code", "dbName", "officialName", "dbCurrency", "officialCurrencyRaw",
            "isoCode", "source", "hedged", "multiCurrency", "confidence"]
    for r_idx, row in enumerate(rows, start=2):
        needs_confirm = row["confidence"] == "需人工確認"
        for c_idx, key in enumerate(keys, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=row[key])
            cell.font = CONFIRM_FONT if needs_confirm else BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(c_idx in (2, 3)))
            if needs_confirm:
                cell.fill = CONFIRM_FILL
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(rows) + 1}"

wb = Workbook()
ws_summary = wb.active
ws_summary.title = "摘要"

ws_tdcc = wb.create_sheet("境外基金 TDCC")
write_sheet(ws_tdcc, data["tdccRows"])
ws_sitca = wb.create_sheet("境內基金 SITCA")
write_sheet(ws_sitca, data["sitcaRows"])

# ---- Summary sheet (formulas reference the two data sheets) ----
ws_summary.sheet_view.showGridLines = False
ws_summary["A1"] = "幣別比對清單摘要（2026-10-21）"
ws_summary["A1"].font = Font(name=FONT, bold=True, size=14)
ws_summary.merge_cells("A1:D1")

ws_summary["A3"] = "資料來源"
ws_summary["B3"] = "工作表"
ws_summary["C3"] = "不一致筆數"
for c in ("A3", "B3", "C3"):
    ws_summary[c].font = HEADER_FONT
    ws_summary[c].fill = HEADER_FILL
    ws_summary[c].border = BORDER
    ws_summary[c].alignment = Alignment(horizontal="center")

tdcc_n = len(data["tdccRows"])
sitca_n = len(data["sitcaRows"])

rows_meta = [
    ("TDCC dataset 3-2（境外基金）", "境外基金 TDCC", tdcc_n),
    ("SITCA CSV（境內基金）", "境內基金 SITCA", sitca_n),
]
r = 4
for source, sheet, n in rows_meta:
    ws_summary.cell(row=r, column=1, value=source).font = BODY_FONT
    ws_summary.cell(row=r, column=2, value=sheet).font = BODY_FONT
    ws_summary.cell(row=r, column=3, value=n).font = BODY_FONT
    for c in range(1, 4):
        ws_summary.cell(row=r, column=c).border = BORDER
    r += 1
total_row = r
ws_summary.cell(row=r, column=1, value="總計").font = Font(name=FONT, bold=True)
ws_summary.cell(row=r, column=3, value=f"=SUM(C4:C{r-1})").font = Font(name=FONT, bold=True)
for c in range(1, 4):
    ws_summary.cell(row=r, column=c).border = BORDER
ws_summary.column_dimensions["A"].width = 30
ws_summary.column_dimensions["B"].width = 20
ws_summary.column_dimensions["C"].width = 14
ws_summary.column_dimensions["D"].width = 14

# ---- 各幣別的檔數（合併兩個工作表的 ISO 代碼欄，F欄） ----
r2 = total_row + 3
ws_summary.cell(row=r2, column=1, value="準備改成的ISO代碼（合併兩表）").font = Font(name=FONT, bold=True, size=12)
r2 += 1
ws_summary.cell(row=r2, column=1, value="ISO代碼").font = HEADER_FONT
ws_summary.cell(row=r2, column=1).fill = HEADER_FILL
ws_summary.cell(row=r2, column=2, value="檔數").font = HEADER_FONT
ws_summary.cell(row=r2, column=2).fill = HEADER_FILL
for c in (1, 2):
    ws_summary.cell(row=r2, column=c).border = BORDER
header_row2 = r2
all_isos = sorted(set(row["isoCode"] for row in data["tdccRows"] + data["sitcaRows"]))
r2 += 1
iso_start_row = r2
for iso in all_isos:
    ws_summary.cell(row=r2, column=1, value=iso).font = BODY_FONT
    formula = (
        f'=COUNTIF(\'境外基金 TDCC\'!$F$2:$F${tdcc_n+1},A{r2})'
        f'+COUNTIF(\'境內基金 SITCA\'!$F$2:$F${sitca_n+1},A{r2})'
    )
    ws_summary.cell(row=r2, column=2, value=formula).font = BODY_FONT
    for c in (1, 2):
        ws_summary.cell(row=r2, column=c).border = BORDER
    r2 += 1
iso_end_row = r2 - 1

# ---- 需人工確認 / 官方幣別為新台幣 ----
r3 = r2 + 2
labels = [
    ("需人工確認的檔數（合併兩表）",
     f"=COUNTIF('境外基金 TDCC'!$J$2:$J${tdcc_n+1},\"需人工確認\")+COUNTIF('境內基金 SITCA'!$J$2:$J${sitca_n+1},\"需人工確認\")"),
    ("官方幣別為新台幣的檔數（合併兩表）",
     f"=COUNTIF('境外基金 TDCC'!$F$2:$F${tdcc_n+1},\"TWD\")+COUNTIF('境內基金 SITCA'!$F$2:$F${sitca_n+1},\"TWD\")"),
]
for label, formula in labels:
    ws_summary.cell(row=r3, column=1, value=label).font = BODY_FONT
    ws_summary.cell(row=r3, column=3, value=formula).font = Font(name=FONT, bold=True)
    for c in (1, 3):
        ws_summary.cell(row=r3, column=c).border = BORDER
    r3 += 1

ws_summary.cell(row=r3 + 1, column=1, value="備註：黃底列＝信心程度「需人工確認」（避險級別、名稱含多種幣別、官方標示新台幣但名稱有外幣字樣，或名稱比對不完全一致）。").font = Font(name=FONT, size=9, italic=True)
ws_summary.merge_cells(start_row=r3 + 1, start_column=1, end_row=r3 + 1, end_column=6)

wb.save(OUT_PATH)
print(f"saved: {OUT_PATH}")
print(f"TDCC rows: {tdcc_n}, SITCA rows: {sitca_n}, total: {tdcc_n + sitca_n}")
